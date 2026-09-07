# -*- coding: utf-8 -*-
"""
QQ群成员导出工具
功能: 选择导出你加入的指定群的成员列表(群昵称 + QQ昵称原名 + QQ号)到 Excel
用法:
  python 导出群成员.py                交互式选择要导出的群
  python 导出群成员.py 1,3,5-8        按列表序号导出
  python 导出群成员.py 关键词1 关键词2  导出群名包含关键词的群
  python 导出群成员.py all            导出全部群
依赖: pip install websockets openpyxl
前置: NapCat 已登录并开启本地 WebSocket 服务 (ws://127.0.0.1:3001)
"""
import asyncio
import json
import re
import sys
from datetime import datetime

for _stream in (sys.stdout, sys.stderr):
    try:
        _stream.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

try:
    import websockets
except ImportError:
    print("缺少依赖 websockets,请先运行: pip install websockets openpyxl")
    sys.exit(1)

WS_URL = "ws://127.0.0.1:3001"
API_TIMEOUT = 60
HEADERS = ["群号", "群名称", "QQ号", "群昵称", "QQ昵称(原名)", "群内角色"]
ROLE_CN = {"owner": "群主", "admin": "管理员", "member": "成员"}


async def call_api(ws, action, params, echo_id):
    await ws.send(json.dumps({"action": action, "params": params, "echo": echo_id}, ensure_ascii=False))
    try:
        while True:
            raw = await asyncio.wait_for(ws.recv(), timeout=API_TIMEOUT)
            msg = json.loads(raw)
            if msg.get("echo") == echo_id:
                return msg
    except asyncio.TimeoutError:
        return {"status": "failed", "retcode": -1, "data": None, "error": "请求超时"}


def get_group_list():
    async def _fetch():
        async with websockets.connect(WS_URL, max_size=64 * 1024 * 1024, open_timeout=15) as ws:
            return await call_api(ws, "get_group_list", {}, "getgroups")

    resp = asyncio.run(_fetch())
    if resp.get("status") != "ok" or resp.get("retcode", -1) != 0:
        print("获取群列表失败:", resp)
        sys.exit(1)
    groups = sorted(resp.get("data") or [], key=lambda g: g.get("group_id", 0))
    return groups


def parse_index_selection(text, count):
    selected = set()
    for part in re.split(r"[,\s;、]+", text.strip()):
        if not part:
            continue
        m = re.fullmatch(r"(\d+)(?:-(\d+))?", part)
        if not m:
            return None
        a, b = int(m.group(1)), int(m.group(2) or m.group(1))
        if a < 1 or b > count or a > b:
            return None
        selected.update(range(a, b + 1))
    return selected if selected else None


def select_groups(groups, argv):
    count = len(groups)
    print("共 %d 个群:" % count)
    for i, g in enumerate(groups, 1):
        name = g.get("group_name") or str(g.get("group_id"))
        members = g.get("member_count")
        extra = " (成员%d)" % members if isinstance(members, int) else ""
        print("  %3d. %s (%s)%s" % (i, name, g.get("group_id"), extra))

    text = ""
    if argv:
        if argv[0].lower() == "all":
            return groups
        text = " ".join(argv)
    else:
        print()
        print("输入要导出的群: 序号(如 1,3,5-8)、关键词(如 智能2404)、或 all 导出全部,多个关键词用空格分隔")
        print("例如: 1,3,5-8  或  智能2404 班委  或  all")
        try:
            text = input("请选择: ").strip()
        except EOFError:
            print("未检测到输入,退出。也可以直接运行: python 导出群成员.py all")
            sys.exit(0)

    if not text:
        print("未选择任何群")
        sys.exit(0)
    if text.lower() == "all":
        return groups

    idxs = parse_index_selection(text, count)
    if idxs is not None:
        picked = [groups[i - 1] for i in sorted(idxs)]
        print("已选择 %d 个群" % len(picked))
        return picked

    keywords = text.split()
    picked = []
    for g in groups:
        name = str(g.get("group_name") or "")
        gid = str(g.get("group_id"))
        if any(k.lower() in name.lower() or k == gid for k in keywords):
            picked.append(g)
    if not picked:
        print("没有匹配的群,请检查关键词")
        sys.exit(0)
    print("关键词匹配到 %d 个群:" % len(picked))
    for g in picked:
        print("  - %s (%s)" % (g.get("group_name") or g.get("group_id"), g.get("group_id")))
    return picked


async def fetch_members(selected):
    async with websockets.connect(WS_URL, max_size=64 * 1024 * 1024, open_timeout=15) as ws:
        seq = 0

        def next_echo():
            nonlocal seq
            seq += 1
            return "seq%d" % seq

        all_rows = []
        per_group = []
        failed = []
        for i, g in enumerate(selected, 1):
            gid = str(g.get("group_id"))
            gname = g.get("group_name") or gid
            print("[%d/%d] %s (%s) " % (i, len(selected), gname, gid), end="", flush=True)
            members = None
            err = None
            for attempt in range(3):
                r = await call_api(ws, "get_group_member_list",
                                   {"group_id": int(gid), "no_cache": False}, next_echo())
                if r.get("status") == "ok" and r.get("retcode") == 0:
                    members = r.get("data") or []
                    break
                err = r.get("error") or r.get("wording") or r
                await asyncio.sleep(2)
            if members is None:
                failed.append((gid, gname, str(err)))
                print("失败,已跳过")
                continue
            rows = []
            for m in members:
                qq = str(m.get("user_id", ""))
                card = m.get("card") or ""
                nickname = m.get("nickname") or ""
                role = ROLE_CN.get(m.get("role"), str(m.get("role", "")))
                rows.append([gid, gname, qq, card, nickname, role])
            all_rows.extend(rows)
            per_group.append((gid, gname, rows))
            print("%d 人" % len(rows))
            await asyncio.sleep(0.6)
        return all_rows, per_group, failed


def sanitize_sheet(name):
    name = re.sub(r"[\x00-\x1f]", "", str(name))
    name = re.sub(r"[\\/?*\[\]:]", "_", name)
    name = name.strip("' ") or "群"
    return name[:31]


def save_excel(path, all_rows, per_group):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "全部成员汇总"
    ws.append(HEADERS)
    for c in range(1, len(HEADERS) + 1):
        ws.cell(row=1, column=c).font = Font(bold=True)
    for row in all_rows:
        ws.append(row)
    for i, w in enumerate([15, 24, 16, 26, 26, 10], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"

    used = set()
    for gid, gname, rows in per_group:
        if not rows:
            continue
        base = sanitize_sheet(gname)
        title = base
        k = 2
        while title in used:
            title = "%s_%d" % (base[:28], k)
            k += 1
        used.add(title)
        s = wb.create_sheet(title)
        s.append(HEADERS)
        for c in range(1, len(HEADERS) + 1):
            s.cell(row=1, column=c).font = Font(bold=True)
        for row in rows:
            s.append(row)
        for i, w in enumerate([15, 24, 16, 26, 26, 10], 1):
            s.column_dimensions[get_column_letter(i)].width = w
        s.freeze_panes = "A2"

    wb.save(path)


def main():
    print("正在连接 NapCat (%s)..." % WS_URL)
    try:
        groups = get_group_list()
    except Exception as e:
        print("运行出错:", e)
        print("若提示连接失败,请先运行 第2步_一键导出.bat (会自动启动并登录NapCat)")
        sys.exit(1)

    selected = select_groups(groups, sys.argv[1:])

    print("开始导出所选群成员...")
    try:
        all_rows, per_group, failed = asyncio.run(fetch_members(selected))
    except Exception as e:
        print("运行出错:", e)
        sys.exit(1)

    if not all_rows:
        print("没有导出到任何数据")
        sys.exit(1)

    out = "群成员导出_%s.xlsx" % datetime.now().strftime("%Y%m%d_%H%M%S")
    try:
        save_excel(out, all_rows, per_group)
    except Exception as e:
        print("写入Excel失败:", e, ",改为输出CSV")
        out = out.replace(".xlsx", ".csv")
        with open(out, "w", encoding="utf-8-sig", newline="") as f:
            f.write(",".join(HEADERS) + "\n")
            for row in all_rows:
                f.write(",".join('"%s"' % str(x).replace('"', '""') for x in row) + "\n")

    print()
    print("导出完成: %s" % out)
    print("共 %d 个群, %d 名成员记录" % (len(per_group), len(all_rows)))
    if failed:
        print("以下群导出失败(已跳过):")
        for gid, gname, err in failed:
            print("  %s (%s): %s" % (gname, gid, err))


if __name__ == "__main__":
    main()
