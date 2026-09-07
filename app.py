# -*- coding: utf-8 -*-
"""
QQ群成员导出工具 - 网页版
扫码登录 -> 选择群 -> 一键导出群成员(群昵称 + QQ昵称原名 + QQ号)到 Excel
本地服务,数据不经过云端
"""
import asyncio
import glob
import json
import os
import re
import socket
import subprocess
import sys
import threading
import time
import uuid
from datetime import datetime

for _s in (sys.stdout, sys.stderr):
    try:
        _s.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

BASE = os.path.dirname(os.path.abspath(__file__))


def _resolve_napcat_dir():
    """定位 NapCat 部署目录: 环境变量 NAPCAT_DIR > 用户目录 NapCatShell > 脚本同目录 NapCatShell"""
    cand = os.environ.get("NAPCAT_DIR")
    if cand and os.path.isdir(cand):
        return cand
    cand = os.path.join(os.path.expanduser("~"), "NapCatShell")
    if os.path.isdir(cand):
        return cand
    cand = os.path.join(BASE, "NapCatShell")
    if os.path.isdir(cand):
        return cand
    raise RuntimeError(
        "未找到 NapCat 部署目录。请先按 README 部署 NapCat,"
        "或通过环境变量 NAPCAT_DIR 指定目录(该目录下需存在 bootmain\\QQ.exe)。")


NAP_DIR = _resolve_napcat_dir()
QQ_EXE = os.path.join(NAP_DIR, "bootmain", "QQ.exe")
QQ_WORKDIR = os.path.join(NAP_DIR, "bootmain")
CFG_DIR = os.path.join(NAP_DIR, "config")
QR_PNG = os.path.join(NAP_DIR, "cache", "qrcode.png")
RUN_LOG = os.path.join(NAP_DIR, "run.log")
EXPORT_DIR = os.path.join(BASE, "exports")
WS_URL = "ws://127.0.0.1:3001"
WEB_PORT = 8123

os.makedirs(EXPORT_DIR, exist_ok=True)

try:
    import psutil
except ImportError:
    subprocess.run([sys.executable, "-m", "pip", "install", "psutil", "--quiet"])
    import psutil

try:
    from flask import Flask, request, jsonify, send_file, Response
except ImportError:
    subprocess.run([sys.executable, "-m", "pip", "install", "flask", "--quiet"])
    from flask import Flask, request, jsonify, send_file, Response

import websockets

app = Flask(__name__)

STATE = {
    "desired": True,      # 是否希望 NapCat 运行
    "state": "offline",   # offline / starting / qr / restarting / ready / error
    "uin": None,
    "nickname": None,
    "message": "",
}
LOCK = threading.Lock()
ROLE_CN = {"owner": "群主", "admin": "管理员", "member": "成员"}
HEADERS = ["群号", "群名称", "QQ号", "群昵称", "QQ昵称(原名)", "群内角色"]


# ---------------- NapCat 进程管理 ----------------

def napcat_running():
    try:
        for p in psutil.process_iter(["name", "exe"]):
            if p.info["name"] == "QQ.exe" and p.info["exe"] and p.info["exe"].lower().startswith(NAP_DIR.lower()):
                return True
    except Exception:
        pass
    return False


def stop_napcat():
    for _ in range(3):
        try:
            killed = False
            for p in psutil.process_iter(["name", "exe"]):
                try:
                    if p.info["name"] == "QQ.exe" and p.info["exe"] and p.info["exe"].lower().startswith(NAP_DIR.lower()):
                        p.kill()
                        killed = True
                except Exception:
                    pass
            if not killed:
                break
        except Exception:
            pass
        time.sleep(1.5)


def start_napcat(uin=None):
    args = [QQ_EXE, "--enable-logging"]
    if uin:
        args += ["-q", str(uin)]
    logf = open(RUN_LOG, "ab", buffering=0)
    subprocess.Popen(args, cwd=QQ_WORKDIR, stdout=logf, stderr=subprocess.STDOUT,
                     creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
                     close_fds=True)


def ws_ready():
    s = socket.socket()
    s.settimeout(1.0)
    try:
        s.connect(("127.0.0.1", 3001))
        return True
    except Exception:
        return False
    finally:
        s.close()


def find_uin():
    for f in glob.glob(os.path.join(CFG_DIR, "onebot11_*.json")):
        m = re.search(r"onebot11_(\d+)\.json", os.path.basename(f))
        if m:
            return m.group(1)
    return None


def find_main_qq():
    """检测用户自己打开的QQ客户端(非NapCat的QQ实例)"""
    try:
        for p in psutil.process_iter(["name", "exe"]):
            if p.info["name"] == "QQ.exe" and p.info["exe"] and not p.info["exe"].lower().startswith(NAP_DIR.lower()):
                return True
    except Exception:
        pass
    return False


def kill_main_qq():
    """关闭用户自己打开的QQ客户端"""
    for _ in range(3):
        try:
            killed = False
            for p in psutil.process_iter(["name", "exe"]):
                try:
                    if p.info["name"] == "QQ.exe" and p.info["exe"] and not p.info["exe"].lower().startswith(NAP_DIR.lower()):
                        p.kill()
                        killed = True
                except Exception:
                    pass
            if not killed:
                break
        except Exception:
            pass
        time.sleep(1.5)


def write_ws_config(uin):
    cfg = {
        "network": {
            "httpServers": [],
            "httpSseServers": [],
            "httpClients": [],
            "websocketServers": [{
                "name": "web_export",
                "enable": True,
                "host": "127.0.0.1",
                "port": 3001,
                "messagePostFormat": "array",
                "reportSelfMessage": False,
                "token": "",
                "enableForcePushEvent": True,
                "debug": False,
                "heartInterval": 30000,
            }],
            "websocketClients": [],
            "plugins": [],
        },
        "musicSignUrl": "",
        "enableLocalFile2Url": False,
        "parseMultMsg": False,
        "imageDownloadProxy": "",
        "timeout": {"baseTimeout": 10000, "uploadSpeedKBps": 256,
                    "downloadSpeedKBps": 256, "maxTimeout": 1800000},
    }
    path = os.path.join(CFG_DIR, "onebot11_%s.json" % uin)
    with open(path, "w", encoding="utf-8") as f:
        json.dump(cfg, f, ensure_ascii=False, indent=2)


# ---------------- OneBot 调用 ----------------

def onebot_call(action, params, timeout=60):
    async def _run():
        async with websockets.connect(WS_URL, max_size=64 * 1024 * 1024, open_timeout=10) as ws:
            await ws.send(json.dumps({"action": action, "params": params, "echo": "req"}, ensure_ascii=False))
            while True:
                raw = await asyncio.wait_for(ws.recv(), timeout=timeout)
                msg = json.loads(raw)
                if msg.get("echo") == "req":
                    return msg
    return asyncio.run(_run())


def get_login_info():
    r = onebot_call("get_login_info", {}, 15)
    if r.get("status") == "ok" and r.get("retcode", -1) == 0 and r.get("data"):
        return r["data"]
    return None


def get_group_list():
    r = onebot_call("get_group_list", {}, 20)
    if r.get("status") != "ok" or r.get("retcode", -1) != 0:
        raise RuntimeError("获取群列表失败: %s" % json.dumps(r, ensure_ascii=False)[:200])
    groups = r.get("data") or []
    groups.sort(key=lambda g: g.get("group_id", 0))
    return groups


def get_group_members(group_id):
    for _ in range(3):
        try:
            r = onebot_call("get_group_member_list", {"group_id": int(group_id), "no_cache": False}, 90)
            if r.get("status") == "ok" and r.get("retcode") == 0:
                return r.get("data") or []
            raise RuntimeError(r.get("wording") or r.get("error") or json.dumps(r, ensure_ascii=False)[:120])
        except Exception as e:
            last = e
            time.sleep(2)
    raise last


# ---------------- 导出 Excel ----------------

def sanitize_sheet(name):
    name = re.sub(r"[\x00-\x1f]", "", str(name))
    name = re.sub(r"[\\/?*\[\]:]", "_", name)
    name = name.strip("' ") or "群"
    return name[:31]


def build_workbook(groups_data):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "全部成员汇总"
    ws.append(HEADERS)
    for c in range(1, len(HEADERS) + 1):
        ws.cell(row=1, column=c).font = Font(bold=True)
    total = 0
    used = set()
    for gid, gname, members in groups_data:
        for m in members:
            qq = str(m.get("user_id", ""))
            card = m.get("card") or ""
            nickname = m.get("nickname") or ""
            role = ROLE_CN.get(m.get("role"), str(m.get("role", "")))
            ws.append([gid, gname, qq, card, nickname, role])
            total += 1
    for i, w in enumerate([15, 26, 16, 28, 28, 10], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"

    for gid, gname, members in groups_data:
        if not members:
            continue
        base = sanitize_sheet(gname)
        title = base
        k = 2
        while title in used:
            title = "%s_%d" % (base[:28], k)
            k += 1
        used.add(title)
        s = wb.create_sheet(title)
        s.append(HEADERS)
        for c in range(1, len(HEADERS) + 1):
            s.cell(row=1, column=c).font = Font(bold=True)
        for m in members:
            qq = str(m.get("user_id", ""))
            card = m.get("card") or ""
            nickname = m.get("nickname") or ""
            role = ROLE_CN.get(m.get("role"), str(m.get("role", "")))
            s.append([gid, gname, qq, card, nickname, role])
        for i, w in enumerate([15, 26, 16, 28, 28, 10], 1):
            s.column_dimensions[get_column_letter(i)].width = w
        s.freeze_panes = "A2"
    return wb, total


def cleanup_exports():
    now = time.time()
    for f in glob.glob(os.path.join(EXPORT_DIR, "*.xlsx")):
        try:
            if now - os.path.getmtime(f) > 3600:
                os.remove(f)
        except Exception:
            pass


# ---------------- 守护线程 ----------------

def supervisor():
    time.sleep(1.5)
    while True:
        try:
            tick()
        except Exception as e:
            with LOCK:
                STATE["message"] = str(e)[:200]
        time.sleep(2)


def tick():
    with LOCK:
        desired = STATE["desired"]
        st = STATE["state"]
    if not desired:
        return

    # 检测用户自己打开的QQ客户端冲突
    if find_main_qq():
        if napcat_running():
            stop_napcat()
        with LOCK:
            STATE["state"] = "conflict"
            STATE["message"] = "检测到你电脑上的QQ客户端正在运行,与NapCat登录冲突"
        return

    uin = find_uin()
    running = napcat_running()

    if uin:
        if st == "ready":
            if not running:
                with LOCK:
                    STATE["state"] = "restarting"
                    STATE["message"] = "NapCat 已断开,正在自动重连..."
                    STATE["restart_at"] = time.time()
                    STATE["restart_attempts"] = 0
            return
        if st == "restarting":
            if ws_ready():
                info = None
                try:
                    info = get_login_info()
                except Exception:
                    info = None
                if info and info.get("user_id"):
                    with LOCK:
                        STATE["state"] = "ready"
                        STATE["uin"] = uin
                        STATE["nickname"] = info.get("nickname")
                        STATE["message"] = "已就绪,可以导出"
                else:
                    with LOCK:
                        STATE["state"] = "qr"
                        STATE["uin_exists"] = True
                        STATE["message"] = "QQ要求手Q验证,已切换为扫码登录,请扫描二维码"
                return
            since = STATE.get("restart_at", 0)
            attempts = STATE.get("restart_attempts", 0)
            if since and time.time() - since > 75:
                with LOCK:
                    STATE["restart_attempts"] = attempts + 1
                    STATE["restart_at"] = time.time()
                if attempts >= 3:
                    with LOCK:
                        STATE["state"] = "error"
                        STATE["message"] = "多次启动失败,请点击「查看日志」或将日志复制给开发者"
                    return
                write_ws_config(uin)
                stop_napcat()
                start_napcat(uin)
                with LOCK:
                    STATE["message"] = "正在重新启动导出服务(第%d次重试)..." % (attempts + 1)
            return
        # 已登录但还没启用服务
        write_ws_config(uin)
        if running:
            stop_napcat()
        with LOCK:
            STATE["state"] = "restarting"
            STATE["uin"] = uin
            STATE["message"] = "正在启动导出服务(约20秒)..."
            STATE["restart_at"] = time.time()
            STATE["restart_attempts"] = 0
        start_napcat(uin)
        return

    # 未登录 -> 扫码流程
    if st in ("offline", "starting", "qr", "error", "conflict"):
        if not running:
            with LOCK:
                STATE["state"] = "qr"
                STATE["uin_exists"] = False
                STATE["message"] = "正在生成二维码..."
            start_napcat(None)
            return
        if st == "qr":
            # 首次扫码登录:登录成功后NapCat会生成onebot11配置,写入WS配置并重启启用
            if not STATE.get("uin_exists"):
                new_uin = find_uin()
                if new_uin:
                    write_ws_config(new_uin)
                    stop_napcat()
                    with LOCK:
                        STATE["state"] = "restarting"
                        STATE["uin"] = new_uin
                        STATE["message"] = "登录成功,正在启用导出服务(约20秒)..."
                        STATE["restart_at"] = time.time()
                        STATE["restart_attempts"] = 0
                    start_napcat(new_uin)
                    return
            # 扫码等待中:检测是否已扫码登录成功
            if ws_ready():
                info = None
                try:
                    info = get_login_info()
                except Exception:
                    info = None
                if info and info.get("user_id"):
                    with LOCK:
                        STATE["state"] = "ready"
                        STATE["uin"] = str(info.get("user_id"))
                        STATE["nickname"] = info.get("nickname")
                        STATE["message"] = "已就绪,可以导出"


# ---------------- HTTP 接口 ----------------

@app.route("/")
def index():
    return Response(HTML_PAGE, mimetype="text/html; charset=utf-8")


@app.route("/api/status")
def api_status():
    with LOCK:
        return jsonify(dict(STATE))


@app.route("/api/login")
def api_login():
    with LOCK:
        STATE["desired"] = True
        if STATE["state"] == "offline":
            STATE["state"] = "starting" if find_uin() else "qr"
            STATE["message"] = "正在启动 NapCat..."
    return jsonify({"ok": True})


@app.route("/api/shutdown", methods=["POST"])
def api_shutdown():
    with LOCK:
        STATE["desired"] = False
        STATE["state"] = "offline"
        STATE["message"] = "已关闭"
    stop_napcat()
    return jsonify({"ok": True})


@app.route("/api/killqq", methods=["POST"])
def api_killqq():
    kill_main_qq()
    with LOCK:
        STATE["state"] = "starting"
        STATE["message"] = "已关闭QQ客户端,正在启动 NapCat..."
    return jsonify({"ok": True})


@app.route("/api/logout", methods=["POST"])
def api_logout():
    """彻底退出登录:停止NapCat并清除本机QQ登录凭证,下次启动需重新扫码"""
    if find_main_qq():
        return jsonify({"error": "请先退出你电脑上打开的QQ客户端,再清除登录状态"}), 409
    stop_napcat()
    for pat in ("onebot11_*.json", "napcat_*.json", "napcat_protocol_*.json"):
        for f in glob.glob(os.path.join(CFG_DIR, pat)):
            try:
                os.remove(f)
            except Exception:
                pass
    auth_dir = os.path.join(os.environ.get("APPDATA", ""), "QQ", "auth")
    for f in glob.glob(os.path.join(auth_dir, "*")):
        try:
            os.remove(f)
        except Exception:
            pass
    with LOCK:
        STATE["desired"] = True
        STATE["state"] = "qr"
        STATE["uin"] = None
        STATE["nickname"] = None
        STATE["message"] = "登录状态已清除,请重新扫码登录"
    return jsonify({"ok": True})


@app.route("/api/qrcode")
def api_qrcode():
    with LOCK:
        st = STATE["state"]
    if st not in ("qr",):
        return jsonify({"error": "not_in_login"}), 404
    if not os.path.exists(QR_PNG):
        return jsonify({"error": "qr_pending"}), 404
    resp = Response(open(QR_PNG, "rb").read(), mimetype="image/png")
    resp.headers["Cache-Control"] = "no-store, no-cache, must-revalidate"
    return resp


@app.route("/api/groups")
def api_groups():
    with LOCK:
        st = STATE["state"]
    if st != "ready":
        return jsonify({"error": "not_ready", "state": st}), 503
    try:
        groups = get_group_list()
    except Exception as e:
        return jsonify({"error": str(e)}), 502
    out = [{"group_id": str(g.get("group_id")), "group_name": g.get("group_name") or "",
            "member_count": g.get("member_count", 0)} for g in groups]
    return jsonify({"groups": out})


@app.route("/api/export", methods=["POST"])
def api_export():
    with LOCK:
        st = STATE["state"]
    if st != "ready":
        return jsonify({"error": "NapCat 未就绪,请先完成登录"}), 503
    data = request.get_json(force=True, silent=True) or {}
    ids = [str(x) for x in (data.get("group_ids") or [])]
    if not ids:
        return jsonify({"error": "请先选择要导出的群"}), 400
    try:
        all_groups = {str(g["group_id"]): g.get("group_name") or "" for g in get_group_list()}
    except Exception as e:
        return jsonify({"error": "获取群列表失败: %s" % e}), 502

    groups_data = []
    failed = []
    for gid in ids:
        try:
            members = get_group_members(gid)
        except Exception as e:
            failed.append({"group_id": gid, "error": str(e)})
            continue
        groups_data.append((gid, all_groups.get(gid, gid), members))
        time.sleep(0.5)

    if not groups_data:
        return jsonify({"error": "所选群均导出失败: %s" % json.dumps(failed, ensure_ascii=False)}), 502

    try:
        wb, total = build_workbook(groups_data)
    except Exception as e:
        return jsonify({"error": "生成Excel失败: %s" % e}), 500

    fname = "群成员导出_%s.xlsx" % datetime.now().strftime("%Y%m%d_%H%M%S")
    path = os.path.join(EXPORT_DIR, fname)
    wb.save(path)
    return jsonify({"ok": True, "url": "/api/download/%s" % fname, "filename": fname,
                    "group_count": len(groups_data), "member_total": total, "failed": failed})


@app.route("/api/download/<path:fname>")
def api_download(fname):
    if ".." in fname or "/" in fname or "\\" in fname:
        return jsonify({"error": "非法文件名"}), 400
    path = os.path.join(EXPORT_DIR, fname)
    if not os.path.exists(path):
        return jsonify({"error": "文件不存在"}), 404
    return send_file(path, as_attachment=True, download_name=fname)


@app.route("/api/log")
def api_log():
    try:
        with open(RUN_LOG, "r", encoding="utf-8", errors="replace") as f:
            lines = f.readlines()[-30:]
        ansi = re.compile(r"\x1b\[[0-9;]*m")
        return jsonify({"log": [ansi.sub("", l).rstrip() for l in lines]})
    except Exception as e:
        return jsonify({"log": ["日志不可用: %s" % e]})


# ---------------- 前端页面 ----------------

HTML_PAGE = r"""<!DOCTYPE html>
<html lang="zh-CN">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>QQ群成员导出工具</title>
<style>
:root{
  --bg:#0d1220; --card:rgba(255,255,255,.045); --border:rgba(255,255,255,.09);
  --text:#e9edf8; --sub:#98a2c3; --accent1:#6c8cff; --accent2:#3ad6ff;
  --green:#3ddc84; --amber:#ffb454; --red:#ff6b6b;
}
*{margin:0;padding:0;box-sizing:border-box}
body{
  min-height:100vh; color:var(--text);
  font-family:"PingFang SC","Microsoft YaHei","Segoe UI",system-ui,sans-serif;
  background:var(--bg);
  background-image:
    radial-gradient(900px 500px at 15% -10%, rgba(108,140,255,.20), transparent 60%),
    radial-gradient(800px 480px at 90% 0%, rgba(58,214,255,.14), transparent 60%),
    radial-gradient(700px 600px at 50% 110%, rgba(120,90,255,.12), transparent 60%);
  background-attachment:fixed;
}
.wrap{max-width:960px;margin:0 auto;padding:36px 20px 80px}
header{display:flex;align-items:center;justify-content:space-between;gap:16px;flex-wrap:wrap;margin-bottom:26px}
.logo{display:flex;align-items:center;gap:14px}
.logo .mark{
  width:46px;height:46px;border-radius:14px;display:flex;align-items:center;justify-content:center;
  background:linear-gradient(135deg,var(--accent1),var(--accent2));font-size:22px;font-weight:800;color:#fff;
  box-shadow:0 6px 24px rgba(108,140,255,.35);
}
.logo h1{font-size:21px;font-weight:700;letter-spacing:.5px}
.logo p{font-size:12.5px;color:var(--sub);margin-top:3px}
.pill{
  display:inline-flex;align-items:center;gap:8px;padding:8px 16px;border-radius:999px;
  background:var(--card);border:1px solid var(--border);backdrop-filter:blur(12px);
  font-size:13.5px;color:var(--sub);transition:.3s;
}
.pill .dot{width:9px;height:9px;border-radius:50%;background:#64748b;flex:none}
.pill.online .dot{background:var(--green);box-shadow:0 0 10px var(--green)}
.pill.qr .dot{background:var(--amber);box-shadow:0 0 10px var(--amber);animation:blink 1.6s infinite}
.pill.work .dot{background:var(--accent2);box-shadow:0 0 10px var(--accent2);animation:blink 1.6s infinite}
.pill.error .dot{background:var(--red)}
@keyframes blink{50%{opacity:.35}}
.card{
  background:var(--card);border:1px solid var(--border);border-radius:18px;padding:28px;
  backdrop-filter:blur(14px);box-shadow:0 12px 40px rgba(0,0,0,.28);margin-bottom:22px;
  animation:fadeUp .5s ease both;
}
@keyframes fadeUp{from{opacity:0;transform:translateY(14px)}to{opacity:1;transform:none}}
.card h2{font-size:16.5px;font-weight:700;display:flex;align-items:center;gap:10px;margin-bottom:18px}
.card h2 .step{
  width:24px;height:24px;border-radius:8px;background:linear-gradient(135deg,var(--accent1),var(--accent2));
  display:inline-flex;align-items:center;justify-content:center;font-size:13px;color:#fff;flex:none;
}
.muted{color:var(--sub);font-size:13px;line-height:1.7}
.qrbox{display:flex;flex-direction:column;align-items:center;gap:16px;padding:6px 0}
.qrframe{
  width:230px;height:230px;background:#fff;border-radius:16px;padding:12px;
  display:flex;align-items:center;justify-content:center;position:relative;
  box-shadow:0 10px 30px rgba(0,0,0,.35);
}
.qrframe img{width:100%;height:100%;object-fit:contain;image-rendering:pixelated}
.qrframe .placeholder{color:#8a93ad;font-size:13px;text-align:center;line-height:1.8}
.qrframe .overlay{
  position:absolute;inset:0;border-radius:16px;display:flex;align-items:center;justify-content:center;
  background:rgba(15,20,35,.72);backdrop-filter:blur(3px);opacity:0;transition:.25s;pointer-events:none;
}
.qrframe.expired .overlay{opacity:1}
.spin{
  width:34px;height:34px;border-radius:50%;border:3px solid rgba(255,255,255,.18);
  border-top-color:var(--accent1);animation:rot .8s linear infinite;
}
@keyframes rot{to{transform:rotate(360deg)}}
.login-ok{display:flex;align-items:center;gap:16px;padding:10px 0}
.avatar{
  width:58px;height:58px;border-radius:50%;flex:none;display:flex;align-items:center;justify-content:center;
  background:linear-gradient(135deg,var(--accent1),var(--accent2));font-size:24px;font-weight:700;color:#fff;
}
.login-ok .who b{font-size:17px}
.login-ok .who span{display:block;color:var(--sub);font-size:13px;margin-top:4px}
.btn{
  border:none;cursor:pointer;font-family:inherit;font-size:14px;font-weight:600;border-radius:12px;
  padding:11px 22px;transition:.22s;color:#fff;
  background:linear-gradient(135deg,var(--accent1),var(--accent2));
  box-shadow:0 6px 20px rgba(108,140,255,.3);
}
.btn:hover{transform:translateY(-2px);box-shadow:0 10px 26px rgba(108,140,255,.42)}
.btn:active{transform:none}
.btn:disabled{opacity:.45;cursor:not-allowed;transform:none;box-shadow:none}
.btn.ghost{background:rgba(255,255,255,.07);border:1px solid var(--border);box-shadow:none;color:var(--text)}
.btn.ghost:hover{border-color:rgba(255,255,255,.25)}
.btn.sm{padding:8px 14px;font-size:12.5px;border-radius:9px}
.btn.danger{background:rgba(255,107,107,.14);border:1px solid rgba(255,107,107,.3);box-shadow:none;color:#ffb3b3}
.btn.danger:hover{border-color:var(--red)}
.searchbox{position:relative;margin-bottom:14px}
.searchbox input{
  width:100%;padding:12px 16px 12px 42px;border-radius:12px;border:1px solid var(--border);
  background:rgba(255,255,255,.05);color:var(--text);font-size:14px;outline:none;transition:.25s;font-family:inherit;
}
.searchbox input:focus{border-color:var(--accent1);background:rgba(255,255,255,.08);box-shadow:0 0 0 3px rgba(108,140,255,.15)}
.searchbox svg{position:absolute;left:14px;top:50%;transform:translateY(-50%);opacity:.5}
.toolbar{display:flex;align-items:center;justify-content:space-between;margin-bottom:10px;gap:10px;flex-wrap:wrap}
.toolbar .left{display:flex;align-items:center;gap:10px;color:var(--sub);font-size:13px}
label.check{display:flex;align-items:center;gap:8px;cursor:pointer;user-select:none}
input[type=checkbox]{width:16px;height:16px;accent-color:var(--accent1);cursor:pointer}
.glist{max-height:420px;overflow-y:auto;border:1px solid var(--border);border-radius:14px;padding:6px}
.glist::-webkit-scrollbar{width:8px}
.glist::-webkit-scrollbar-thumb{background:rgba(255,255,255,.14);border-radius:99px}
.grow{
  display:flex;align-items:center;gap:12px;padding:11px 12px;border-radius:10px;cursor:pointer;transition:.18s;
}
.grow:hover{background:rgba(255,255,255,.06)}
.grow .info{flex:1;min-width:0}
.grow .name{font-size:14px;white-space:nowrap;overflow:hidden;text-overflow:ellipsis}
.grow .gid{font-size:12px;color:var(--sub);margin-top:2px;font-family:Consolas,monospace}
.badge{
  flex:none;font-size:11.5px;color:var(--accent2);background:rgba(58,214,255,.1);
  border:1px solid rgba(58,214,255,.25);padding:3px 10px;border-radius:999px;
}
.empty{text-align:center;color:var(--sub);padding:34px 0;font-size:13.5px}
.footbar{display:flex;align-items:center;justify-content:space-between;gap:14px;margin-top:16px;flex-wrap:wrap}
.sel{
  font-size:13.5px;color:var(--sub);
}
.sel b{color:var(--accent2);font-size:15px}
.result{text-align:center;padding:10px 0 4px}
.result .big{font-size:30px;margin-bottom:12px}
.result .stats{display:flex;justify-content:center;gap:26px;margin:18px 0;flex-wrap:wrap}
.result .stats div{text-align:center}
.result .stats b{display:block;font-size:22px;background:linear-gradient(135deg,var(--accent1),var(--accent2));-webkit-background-clip:text;background-clip:text;color:transparent}
.result .stats span{font-size:12.5px;color:var(--sub)}
.failnote{background:rgba(255,180,84,.08);border:1px solid rgba(255,180,84,.25);border-radius:10px;padding:10px 14px;margin-top:14px;font-size:12.5px;color:#ffd9a8;text-align:left}
footer{margin-top:30px;display:flex;align-items:center;justify-content:space-between;gap:14px;flex-wrap:wrap}
.hint{color:var(--sub);font-size:12px;line-height:1.8;max-width:640px}
.toast{
  position:fixed;left:50%;bottom:34px;transform:translateX(-50%) translateY(20px);opacity:0;
  background:#1c2438;border:1px solid var(--border);border-radius:12px;padding:12px 20px;
  font-size:13.5px;box-shadow:0 10px 30px rgba(0,0,0,.4);transition:.3s;z-index:99;max-width:88vw;
}
.toast.show{opacity:1;transform:translateX(-50%) translateY(0)}
.toast.err{border-color:rgba(255,107,107,.5);color:#ffb3b3}
.toast.ok{border-color:rgba(61,220,132,.5);color:#b9f4d3}
.hidden{display:none!important}
@media (max-width:640px){ .wrap{padding:24px 12px 60px} .card{padding:20px} }
</style>
</head>
<body>
<div class="wrap">
  <header>
    <div class="logo">
      <div class="mark">Q</div>
      <div>
        <h1>QQ群成员导出工具</h1>
        <p>群昵称 &amp; QQ原名 一一对应 · 本地运行 · 数据不经云端</p>
      </div>
    </div>
    <div class="pill" id="pill"><span class="dot"></span><span id="pillText">未启动</span></div>
  </header>

  <div class="card hidden" id="cardLogin">
    <h2><span class="step">1</span> 登录 QQ</h2>
    <div id="loginQr" class="hidden">
      <div class="qrbox">
        <div class="qrframe" id="qrframe">
          <div class="placeholder" id="qrPlaceholder">二维码生成中…</div>
          <img id="qrImg" class="hidden" alt="登录二维码">
          <div class="overlay"><div class="spin"></div></div>
        </div>
        <div class="muted">
          请使用<b>手机QQ「扫一扫」</b>扫描二维码,并在手机上确认登录<br>
          二维码约2分钟自动刷新,请尽快扫码
        </div>
        <div class="muted" id="qrMsg" style="color:#ffd9a8"></div>
        <button class="btn ghost sm" id="btnRefreshQr" onclick="refreshQr()">刷新二维码</button>
      </div>
    </div>
    <div id="loginWorking" class="hidden">
      <div style="display:flex;align-items:center;gap:14px;padding:8px 0">
        <div class="spin" id="workSpin"></div>
        <div style="flex:1"><b id="workTitle">正在启动登录…</b><div class="muted" id="workMsg"></div></div>
        <button class="btn sm hidden" id="btnStart" onclick="apiStart()">启动 NapCat</button>
      </div>
    </div>
    <div id="loginOk" class="hidden">
      <div class="login-ok">
        <div class="avatar" id="avatarTxt">Q</div>
        <div class="who">
          <b id="nickTxt">-</b>
          <span id="uinTxt">QQ号 -</span>
        </div>
        <button class="btn ghost sm" style="margin-left:auto" onclick="apiLogout()">退出登录(重新扫码)</button>
      </div>
      <div class="muted" style="margin-top:14px">✓ 登录成功,可在下方选择群并导出成员</div>
    </div>
    <div id="loginConflict" class="hidden">
      <div style="display:flex;align-items:flex-start;gap:14px;padding:8px 0">
        <div style="font-size:30px;flex:none">⚠️</div>
        <div style="flex:1">
          <b>检测到你电脑上的 QQ 客户端正在运行</b>
          <div class="muted" style="margin-top:8px">
            同一账号无法同时登录,请先退出你电脑上正在运行的 QQ 客户端
            (系统托盘 QQ 图标右键 → 退出),或点击下方按钮自动关闭:<br><br>
            <span style="color:#ffd9a8">注意:自动关闭会直接结束你的 QQ 聊天窗口,导出完成后可重新打开 QQ。</span>
          </div>
          <div style="margin-top:14px;display:flex;gap:10px;flex-wrap:wrap">
            <button class="btn sm" onclick="killMainQQ()">自动关闭 QQ 并继续</button>
            <button class="btn ghost sm" onclick="location.reload()">我已手动关闭,刷新状态</button>
          </div>
        </div>
      </div>
    </div>
  </div>

  <div class="card hidden" id="cardMain">
    <h2><span class="step">2</span> 选择要导出的群</h2>
    <div class="searchbox">
      <svg width="17" height="17" viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2.2" stroke-linecap="round"><circle cx="11" cy="11" r="7"/><path d="m21 21-4.3-4.3"/></svg>
      <input type="text" id="searchInput" placeholder="输入群号或群名搜索,例如: 智能2404 或 787365057" oninput="renderGroups()">
    </div>
    <div class="toolbar">
      <div class="left">
        <label class="check"><input type="checkbox" id="chkAll" onchange="toggleAll()"> 全选</label>
        <span id="shownCount"></span>
      </div>
      <span class="sel">已选 <b id="selCount">0</b> 个群</span>
    </div>
    <div class="glist" id="glist"><div class="empty">加载中…</div></div>
    <div class="footbar">
      <span class="muted">群昵称为空的成员表示未设置群昵称</span>
      <button class="btn" id="btnExport" onclick="doExport()" disabled>导出所选群成员</button>
    </div>
  </div>

  <div class="card hidden" id="cardResult">
    <h2><span class="step">3</span> 导出完成</h2>
    <div class="result">
      <div class="big">🎉</div>
      <div class="stats">
        <div><b id="rsGroups">0</b><span>导出群数</span></div>
        <div><b id="rsMembers">0</b><span>成员记录</span></div>
      </div>
      <button class="btn" id="btnDownload" onclick="doDownload()">下载 Excel 文件</button>
      <div class="muted" style="margin-top:10px" id="rsFile"></div>
      <div class="failnote hidden" id="rsFailed"></div>
      <div style="margin-top:18px"><button class="btn ghost sm" onclick="resetResult()">继续导出其他群</button></div>
    </div>
  </div>

  <footer>
    <div class="hint">
      提示:本工具基于开源框架 NapCatQQ,通过官方QQ客户端扫码登录,不读取你的密码。<br>
      登录期间,你电脑上原本的QQ客户端可能掉线;导出完成后重新登录QQ即可。第三方框架存在极小风控风险,请用完即关,不要长期挂机。
    </div>
    <div style="display:flex;gap:10px">
      <button class="btn ghost sm" onclick="viewLog()">查看日志</button>
      <button class="btn danger sm" onclick="apiShutdown()">关闭 NapCat</button>
    </div>
  </footer>
</div>
<div class="toast" id="toast"></div>

<script>
var STATUS = {state:'offline'};
var GROUPS = [];
var selected = new Set();
var EXPORTING = false;

function toast(msg, type){
  var t = document.getElementById('toast');
  t.textContent = msg;
  t.className = 'toast show' + (type ? ' ' + type : '');
  clearTimeout(t._h);
  t._h = setTimeout(function(){ t.className = 'toast'; }, 3600);
}

function api(path, opt){
  return fetch(path, opt).then(function(r){
    return r.json().catch(function(){ return {error:'服务异常'} }).then(function(j){
      if(!r.ok){ j._status = r.status; }
      return j;
    });
  });
}

function renderStatus(){
  var pill = document.getElementById('pill');
  var txt = document.getElementById('pillText');
  var cLogin = document.getElementById('cardLogin');
  var cMain = document.getElementById('cardMain');
  var q = document.getElementById('loginQr');
  var w = document.getElementById('loginWorking');
  var ok = document.getElementById('loginOk');
  var cf = document.getElementById('loginConflict');
  pill.className = 'pill';
  var s = STATUS.state;
  if(s === 'ready'){
    pill.classList.add('online'); txt.textContent = '已登录在线';
    cLogin.classList.remove('hidden'); cMain.classList.remove('hidden');
    q.classList.add('hidden'); w.classList.add('hidden'); ok.classList.remove('hidden'); cf.classList.add('hidden');
    document.getElementById('uinTxt').textContent = 'QQ号 ' + (STATUS.uin || '-');
    document.getElementById('nickTxt').textContent = STATUS.nickname || '已登录';
    document.getElementById('avatarTxt').textContent = (STATUS.nickname || 'Q').slice(0,1);
  } else if(s === 'conflict'){
    pill.classList.add('error'); txt.textContent = '检测到QQ冲突';
    cLogin.classList.remove('hidden'); cMain.classList.add('hidden');
    q.classList.add('hidden'); w.classList.add('hidden'); ok.classList.add('hidden'); cf.classList.remove('hidden');
  } else if(s === 'qr'){
    pill.classList.add('qr'); txt.textContent = '等待扫码登录';
    cLogin.classList.remove('hidden'); cMain.classList.add('hidden');
    q.classList.remove('hidden'); w.classList.add('hidden'); ok.classList.add('hidden'); cf.classList.add('hidden');
    document.getElementById('qrMsg').textContent = STATUS.message || '';
    loadQr();
  } else if(s === 'restarting' || s === 'starting'){
    pill.classList.add('work'); txt.textContent = s === 'starting' ? '正在启动' : '正在启用导出服务';
    cLogin.classList.remove('hidden'); cMain.classList.add('hidden');
    q.classList.add('hidden'); w.classList.remove('hidden'); ok.classList.add('hidden'); cf.classList.add('hidden');
    document.getElementById('workSpin').classList.remove('hidden');
    document.getElementById('btnStart').classList.add('hidden');
    document.getElementById('workTitle').textContent = s === 'starting' ? '正在启动 NapCat…' : '登录成功,正在启用导出服务…';
    document.getElementById('workMsg').textContent = STATUS.message || '大约需要 20-40 秒,请稍候';
  } else if(s === 'error'){
    pill.classList.add('error'); txt.textContent = '异常';
    cLogin.classList.remove('hidden'); cMain.classList.add('hidden');
    q.classList.add('hidden'); w.classList.remove('hidden'); ok.classList.add('hidden'); cf.classList.add('hidden');
    document.getElementById('workSpin').classList.remove('hidden');
    document.getElementById('btnStart').classList.add('hidden');
    document.getElementById('workTitle').textContent = '启动异常';
    document.getElementById('workMsg').textContent = STATUS.message || '';
  } else {
    txt.textContent = '未运行';
    cLogin.classList.remove('hidden'); cMain.classList.add('hidden');
    q.classList.add('hidden'); w.classList.remove('hidden'); ok.classList.add('hidden'); cf.classList.add('hidden');
    document.getElementById('workSpin').classList.add('hidden');
    document.getElementById('btnStart').classList.remove('hidden');
    document.getElementById('workTitle').textContent = 'NapCat 未运行';
    document.getElementById('workMsg').textContent = STATUS.message || '点击右侧按钮启动并登录';
  }
  var btn = document.getElementById('btnExport');
  if(btn) btn.disabled = selected.size === 0 || EXPORTING;
  document.getElementById('selCount').textContent = selected.size;
}

function refreshQr(){
  var img = document.getElementById('qrImg');
  var frame = document.getElementById('qrframe');
  img.classList.add('hidden');
  frame.classList.remove('expired');
  loadQr(true);
}

function loadQr(force){
  var img = document.getElementById('qrImg');
  var ph = document.getElementById('qrPlaceholder');
  var frame = document.getElementById('qrframe');
  var src = '/api/qrcode?t=' + Date.now();
  if(!force && img.dataset.src === src) return;
  img.onload = function(){
    ph.classList.add('hidden'); img.classList.remove('hidden');
    img.dataset.loadedAt = Date.now();
    frame.classList.remove('expired');
  };
  img.onerror = function(){
    img.classList.add('hidden'); ph.classList.remove('hidden');
    ph.textContent = '二维码生成中…';
  };
  img.dataset.src = src;
  img.src = src;
}

function loadGroups(){
  if(STATUS.state !== 'ready') return;
  api('/api/groups').then(function(j){
    if(j.error){ toast(j.error || '获取群列表失败','err'); return; }
    GROUPS = j.groups || [];
    selected.clear();
    var ids = new Set(Array.from(GROUPS).map(function(g){ return String(g.group_id); }));
    var selectedArr = Array.from(selected).filter(function(x){ return ids.has(x); });
    selected = new Set(selectedArr);
    document.getElementById('chkAll').checked = false;
    renderGroups();
    renderStatus();
  }).catch(function(){ toast('获取群列表失败','err'); });
}

function renderGroups(){
  var kw = (document.getElementById('searchInput').value || '').trim().toLowerCase();
  var list = document.getElementById('glist');
  var arr = GROUPS.filter(function(g){
    if(!kw) return true;
    return (g.group_name || '').toLowerCase().indexOf(kw) >= 0 || String(g.group_id).indexOf(kw) >= 0;
  });
  document.getElementById('shownCount').textContent = '共 ' + arr.length + ' 个群';
  if(!arr.length){ list.innerHTML = '<div class="empty">没有匹配的群</div>'; return; }
  list.innerHTML = arr.map(function(g){
    var gid = String(g.group_id);
    var ck = selected.has(gid) ? 'checked' : '';
    return '<div class="grow" onclick="toggleRow(this)" data-gid="' + gid + '">' +
      '<input type="checkbox" ' + ck + ' data-gid="' + gid + '" onclick="event.stopPropagation();toggleRow(null,this)">' +
      '<div class="info"><div class="name"></div><div class="gid">' + gid + '</div></div>' +
      '<span class="badge">' + (g.member_count || 0) + ' 人</span></div>';
  }).join('');
  list.querySelectorAll('.grow').forEach(function(row){
    var gid = row.getAttribute('data-gid');
    var g = GROUPS.find(function(x){ return String(x.group_id) === gid; });
    if(g) row.querySelector('.name').textContent = g.group_name || gid;
  });
}

function toggleRow(row, cb){
  var gid = cb ? cb.getAttribute('data-gid') : row.getAttribute('data-gid');
  if(cb ? cb.checked : !selected.has(gid)){
    selected.add(gid);
    if(cb) cb.checked = true; else { var c = row.querySelector('input'); if(c) c.checked = true; }
  } else {
    selected.delete(gid);
    if(cb) cb.checked = false; else { var c = row.querySelector('input'); if(c) c.checked = false; }
  }
  document.getElementById('chkAll').checked = selected.size === GROUPS.length && GROUPS.length > 0;
  renderStatus();
}

function toggleAll(){
  var on = document.getElementById('chkAll').checked;
  var kw = (document.getElementById('searchInput').value || '').trim().toLowerCase();
  var arr = GROUPS.filter(function(g){
    if(!kw) return true;
    return (g.group_name || '').toLowerCase().indexOf(kw) >= 0 || String(g.group_id).indexOf(kw) >= 0;
  });
  if(on){ arr.forEach(function(g){ selected.add(String(g.group_id)); }); }
  else { arr.forEach(function(g){ selected.delete(String(g.group_id)); }); }
  renderGroups();
  renderStatus();
}

function doExport(){
  if(!selected.size || EXPORTING) return;
  EXPORTING = true;
  var btn = document.getElementById('btnExport');
  btn.disabled = true;
  btn.textContent = '正在导出,请稍候…';
  toast('正在导出 ' + selected.size + ' 个群的成员(大群可能需要 1-2 分钟)', 'ok');
  api('/api/export', {
    method:'POST',
    headers:{'Content-Type':'application/json'},
    body: JSON.stringify({group_ids: Array.from(selected)})
  }).then(function(j){
    EXPORTING = false;
    btn.textContent = '导出所选群成员';
    renderStatus();
    if(!j.ok && j.error){ toast(j.error, 'err'); return; }
    window.LAST_RESULT = j;
    var c = document.getElementById('cardResult');
    c.classList.remove('hidden');
    document.getElementById('rsGroups').textContent = j.group_count;
    document.getElementById('rsMembers').textContent = j.member_total;
    document.getElementById('rsFile').textContent = j.filename;
    var f = document.getElementById('rsFailed');
    if(j.failed && j.failed.length){
      f.classList.remove('hidden');
      f.innerHTML = '以下群导出失败:<br>' + j.failed.map(function(x){
        return '· ' + x.group_id + ' — ' + x.error;
      }).join('<br>');
    } else { f.classList.add('hidden'); }
    c.scrollIntoView({behavior:'smooth'});
    doDownload();
  }).catch(function(e){
    EXPORTING = false;
    btn.textContent = '导出所选群成员';
    renderStatus();
    toast('导出失败: ' + e, 'err');
  });
}

function doDownload(){
  if(!window.LAST_RESULT) return;
  var a = document.createElement('a');
  a.href = window.LAST_RESULT.url;
  a.download = window.LAST_RESULT.filename;
  document.body.appendChild(a);
  a.click();
  a.remove();
}

function resetResult(){
  document.getElementById('cardResult').classList.add('hidden');
  window.LAST_RESULT = null;
}

function apiStart(){
  api('/api/login').then(function(){
    toast('正在启动 NapCat...','ok');
  });
}

function apiShutdown(){
  if(!confirm('确定关闭 NapCat 吗?关闭后需重新扫码或启动登录。')) return;
  api('/api/shutdown', {method:'POST'}).then(function(){
    toast('NapCat 已关闭','ok');
    document.getElementById('cardMain').classList.add('hidden');
    document.getElementById('cardResult').classList.add('hidden');
  });
}

function killMainQQ(){
  if(!confirm('将自动关闭你电脑上正在运行的QQ客户端,继续吗?')) return;
  api('/api/killqq', {method:'POST'}).then(function(){
    toast('已关闭QQ客户端,正在启动登录...','ok');
  });
}

function apiLogout(){
  if(!confirm('将清除本机的QQ登录状态(不影响手机QQ)。\n\n之后再次登录需重新扫码,且你电脑上的QQ客户端下次打开时也需要重新登录。继续吗?')) return;
  api('/api/logout', {method:'POST'}).then(function(j){
    if(j.error){ toast(j.error, 'err'); return; }
    toast('已清除登录状态,请重新扫码','ok');
    document.getElementById('cardMain').classList.add('hidden');
    document.getElementById('cardResult').classList.add('hidden');
    GROUPS = [];
    selected = new Set();
  });
}

function viewLog(){
  api('/api/log').then(function(j){
    var lines = (j.log || []).join('\n');
    toast('日志已复制到剪贴板(共 ' + (j.log||[]).length + ' 行)');
    if(navigator.clipboard){ navigator.clipboard.writeText(lines); }
    console.log(lines);
  });
}

function poll(){
  api('/api/status').then(function(j){
    var was = STATUS.state;
    STATUS = j;
    if(was !== 'ready' && STATUS.state === 'ready'){ loadGroups(); }
    if(STATUS.state === 'ready' && !GROUPS.length){ loadGroups(); }
    renderStatus();
  }).catch(function(){ });
}
setInterval(poll, 2000);
setInterval(function(){
  if(STATUS.state === 'qr'){ loadQr(true); }
}, 4000);
poll();
</script>
</body>
</html>
"""

if __name__ == "__main__":
    cleanup_exports()
    threading.Thread(target=supervisor, daemon=True).start()
    print("=" * 52)
    print("  QQ群成员导出工具(网页版)已启动")
    print("  请在浏览器打开: http://127.0.0.1:%d" % WEB_PORT)
    print("  关闭本窗口即停止服务")
    print("=" * 52)
    try:
        app.run(host="127.0.0.1", port=WEB_PORT, threaded=True, debug=False)
    except OSError as e:
        print("端口 %d 被占用: %s" % (WEB_PORT, e))
        print("如果网页版已经在运行,直接使用浏览器打开即可")
        sys.exit(1)
